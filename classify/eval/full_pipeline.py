"""Full framework pipeline run: C1 + C2 + C3 + C4 (methodology Sec 3), all
ablation levels, with Markdown + Excel + chart outputs.

  C1 — compound-aware module: backbone_only vs +C1_detection_only vs +C1
       (full detect+decompress) — the complete Sec 9.2 ablation ladder.
  C2 — defensible head: temperature calibration (ECE, reliability bins) +
       Integrated-Gradients deletion-curve faithfulness check.
  C3 — evaluation protocol: per-parent-file splits, both fragment sizes
       separately, per-class metrics, multi-seed paired t-test, cost.
  C4 — open-set: MSP abstention vs OOD fragments (pdf/jpeg/random),
       AUROC + rejection sweep + confidence distributions.

Hyperparameters: read from results/new_result/hypertune.json (best row by
min-class-F1) if present, else fall back to the lr=3e-4 recipe. Scale is
moderate, not publication scale.
"""

import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import torch
from torch.utils.data import DataLoader

from classify.baselines.raw_byte_transformer import RawByteTransformerBaseline
from classify.data.custom_corpus import build_dataset
from classify.data.ood_corpus import OODFragmentDataset
from classify.eval.ablation_runner import run_ablation, to_markdown_table
from classify.eval.deletion_curve import deletion_accuracy_curve
from classify.eval.flops import estimate_transformer_flops
from classify.eval.metrics import inference_throughput, param_count
from classify.modules.c1_compound.dual_view_model import C1Classifier
from classify.modules.c1_compound.train_refiner import train_decompression_confidence_cnn
from classify.modules.heads.defensible_head import TemperatureScaler, expected_calibration_error
from classify.modules.heads.openset_head import auroc_id_vs_ood, msp_scores, rejection_curve
from classify.training.dataset import LABELS, FragmentDataset
from classify.training.train import inverse_frequency_weights, train_classifier

OUT_DIR = "results/new_result"
SEEDS = [0, 1, 2]
DELETION_N_PER_CLASS = 25
DELETION_FRAGMENT_SIZE = 512
CONFIG_NAMES = ("backbone_only", "+C1_detection_only", "+C1")

DEFAULT_TUNING = {"lr": 3e-4, "balance": "oversample", "epochs": 8}

RUN_CONFIG = {
    512: dict(n_per_size=30, model_kwargs=dict(d_model=128, n_heads=8, n_layers=4, d_ff=512),
              batch_size=32, refiner_epochs=3),
    4096: dict(n_per_size=10, model_kwargs=dict(d_model=96, n_heads=4, n_layers=3, d_ff=192),
               batch_size=16, refiner_epochs=2),
}


def load_tuning() -> dict:
    path = os.path.join(OUT_DIR, "hypertune.json")
    if os.path.isfile(path):
        with open(path) as f:
            rows = json.load(f)
        if rows:
            best = rows[0]  # pre-sorted by min-class-F1
            print(f"Using tuned hyperparameters: {best}", flush=True)
            return best
    print(f"hypertune.json not found — using defaults {DEFAULT_TUNING}", flush=True)
    return DEFAULT_TUNING


def _balance_flags(tuning: dict):
    oversample = tuning["balance"] == "oversample"
    return oversample, not oversample  # (oversample, use_class_weights)


def build_model(config_name: str, fragment_size: int, refiner):
    kw = RUN_CONFIG[fragment_size]["model_kwargs"]
    if config_name == "backbone_only":
        return RawByteTransformerBaseline(n_classes=len(LABELS), max_len=fragment_size, **kw)
    detection_only = config_name == "+C1_detection_only"
    return C1Classifier(n_classes=len(LABELS), max_len=fragment_size,
                         pretrained_refiner=refiner, detection_only=detection_only, **kw)


def train_analysis_model(config_name: str, corpus, fragment_size: int, device: str,
                          refiner, tuning: dict):
    cfg = RUN_CONFIG[fragment_size]
    oversample, use_cw = _balance_flags(tuning)
    train_ds = FragmentDataset(corpus["train"], fragment_size=fragment_size, oversample=oversample)
    val_ds = FragmentDataset(corpus["val"], fragment_size=fragment_size)
    class_weights = inverse_frequency_weights(train_ds, n_classes=len(LABELS)) if use_cw else None

    model = build_model(config_name, fragment_size, refiner)
    train_classifier(model, train_ds, val_ds, epochs=tuning["epochs"], batch_size=cfg["batch_size"],
                      lr=tuning["lr"], seed=0, device=device, class_weights=class_weights)
    return model


def reliability_bins(conf: np.ndarray, correct: np.ndarray, n_bins: int = 10) -> list:
    edges = np.linspace(0, 1, n_bins + 1)
    bins = []
    for i in range(n_bins):
        mask = (conf > edges[i]) & (conf <= edges[i + 1]) if i > 0 else (conf >= edges[i]) & (conf <= edges[i + 1])
        if mask.sum() == 0:
            continue
        bins.append({"conf": float(conf[mask].mean()), "acc": float(correct[mask].mean()),
                      "count": int(mask.sum())})
    return bins


def c2_analysis(model, corpus, fragment_size: int, device: str) -> dict:
    val_ds = FragmentDataset(corpus["val"], fragment_size=fragment_size)
    loader = DataLoader(val_ds, batch_size=64)
    model.eval()
    all_logits, all_labels = [], []
    with torch.no_grad():
        for x, y in loader:
            all_logits.append(model(x.to(device)))
            all_labels.append(y.to(device))
    logits, labels = torch.cat(all_logits), torch.cat(all_labels)

    probs = torch.softmax(logits, dim=1)
    conf_before, pred = probs.max(dim=1)
    correct = (pred == labels)
    ece_before = expected_calibration_error(conf_before, correct)

    scaler = TemperatureScaler().to(device)
    scaler.fit(logits, labels)
    conf_after, _ = torch.softmax(scaler(logits), dim=1).max(dim=1)
    ece_after = expected_calibration_error(conf_after, correct)

    conf_b = conf_before.detach().cpu().numpy()
    conf_a = conf_after.detach().cpu().numpy()
    corr = correct.detach().cpu().numpy()
    return {"ece_before": ece_before, "ece_after": ece_after,
            "temperature": scaler.log_temperature.exp().item(),
            "reliability_before": reliability_bins(conf_b, corr),
            "reliability_after": reliability_bins(conf_a, corr),
            "scaler": scaler}


def c4_analysis(model, corpus, fragment_size: int, device: str, scaler) -> dict:
    test_ds = FragmentDataset(corpus["test"], fragment_size=fragment_size)
    ood_ds = OODFragmentDataset(fragment_size=fragment_size)

    id_conf, id_pred, id_labels = msp_scores(model, test_ds, device, temperature_scaler=scaler)
    ood_conf, _, _ = msp_scores(model, ood_ds, device, temperature_scaler=scaler)

    return {
        "n_id": len(id_conf), "n_ood": len(ood_conf),
        "auroc": auroc_id_vs_ood(id_conf, ood_conf),
        "rejection_curve": rejection_curve(id_conf, id_pred, id_labels, ood_conf),
        "id_conf_mean": float(id_conf.mean()), "ood_conf_mean": float(ood_conf.mean()),
        "id_conf": id_conf.tolist(), "ood_conf": ood_conf.tolist(),
    }


def cost_analysis(model, corpus, fragment_size: int, device: str) -> dict:
    kw = RUN_CONFIG[fragment_size]["model_kwargs"]
    val_ds = FragmentDataset(corpus["val"], fragment_size=fragment_size)
    return {
        "n_params": param_count(model),
        "flops_per_fragment": estimate_transformer_flops(
            d_model=kw["d_model"], n_layers=kw["n_layers"], d_ff=kw["d_ff"], seq_len=fragment_size),
        "throughput_frags_per_sec": inference_throughput(
            model, val_ds, device, batch_size=32)["fragments_per_sec"],
    }


def _corpus_stats(corpus, fragment_size: int) -> list:
    rows = []
    for split in ("train", "val", "test"):
        frags = [f for f in corpus[split] if f["fragment_size"] == fragment_size]
        rows.append({"split": split,
                      "doc": sum(1 for f in frags if f["label"] == "doc"),
                      "docx": sum(1 for f in frags if f["label"] == "docx")})
    return rows


def run_size(fragment_size: int, device: str, tuning: dict) -> dict:
    cfg = RUN_CONFIG[fragment_size]
    corpus = build_dataset(seed=0, n_per_size=cfg["n_per_size"], sizes=(fragment_size,))
    oversample, use_cw = _balance_flags(tuning)

    print(f"=== {fragment_size}B: C1 ablation (3 levels x {len(SEEDS)} seeds) ===", flush=True)
    ablation = run_ablation(
        corpus=corpus, fragment_size=fragment_size, seeds=SEEDS,
        model_kwargs=cfg["model_kwargs"], epochs=tuning["epochs"],
        batch_size=cfg["batch_size"], lr=tuning["lr"], device=device, out_dir=OUT_DIR,
        refiner_epochs=cfg["refiner_epochs"], oversample=oversample,
        use_class_weights=use_cw, include_detection_only=True,
    )
    print(to_markdown_table(ablation), flush=True)

    refiner = train_decompression_confidence_cnn(
        corpus["train"], fragment_size=fragment_size, epochs=cfg["refiner_epochs"],
        batch_size=cfg["batch_size"], device=device,
    )

    out = {"ablation": ablation, "corpus_stats": _corpus_stats(corpus, fragment_size),
           "c2": {}, "c4": {}, "cost": {}, "deletion_curve": None, "tuning": tuning}

    for config_name in CONFIG_NAMES:
        print(f"=== {fragment_size}B: C2/C4 analysis ({config_name}) ===", flush=True)
        model = train_analysis_model(config_name, corpus, fragment_size, device, refiner, tuning)

        c2 = c2_analysis(model, corpus, fragment_size, device)
        out["c4"][config_name] = c4_analysis(model, corpus, fragment_size, device, c2["scaler"])
        out["cost"][config_name] = cost_analysis(model, corpus, fragment_size, device)
        out["c2"][config_name] = {k: v for k, v in c2.items() if k != "scaler"}

        if fragment_size == DELETION_FRAGMENT_SIZE and config_name == "+C1":
            print(f"=== {fragment_size}B: C2 deletion curve ===", flush=True)
            test_ds = FragmentDataset(corpus["test"], fragment_size=fragment_size)
            out["deletion_curve"] = deletion_accuracy_curve(
                model, test_ds, LABELS, device, n_per_class=DELETION_N_PER_CLASS)

        del model
        if device == "cuda":
            torch.cuda.empty_cache()
    return out


# ---------------------------------------------------------------- charts ---

def make_charts(results: dict, out_dir: str):
    # 1. macro-F1 grouped bar across ablation levels and sizes
    fig, ax = plt.subplots(figsize=(8, 4.5))
    sizes = list(results.keys())
    x = np.arange(len(sizes))
    width = 0.25
    for i, config in enumerate(CONFIG_NAMES):
        means = [results[fs]["ablation"]["configs"][config]["macro_f1_mean"] for fs in sizes]
        stds = [results[fs]["ablation"]["configs"][config]["macro_f1_std"] for fs in sizes]
        ax.bar(x + i * width, means, width, yerr=stds, label=config, capsize=4)
    ax.set_xticks(x + width)
    ax.set_xticklabels([f"{fs}B" for fs in sizes])
    ax.set_ylabel("macro-F1 (mean ± std, 3 seeds)")
    ax.set_title("Ablation levels by fragment size")
    ax.legend()
    fig.tight_layout()
    fig.savefig(os.path.join(out_dir, "chart_macro_f1_ablation.png"), dpi=150)
    plt.close(fig)

    # 2. per-class F1 grouped bar
    fig, axes = plt.subplots(1, len(sizes), figsize=(7 * len(sizes), 4.5))
    axes = np.atleast_1d(axes)
    for ax, fs in zip(axes, sizes):
        x = np.arange(len(CONFIG_NAMES))
        width = 0.35
        for j, cls in enumerate(LABELS):
            vals = [np.mean([r["per_class"][cls]["f1"]
                              for r in results[fs]["ablation"]["configs"][c]["runs"]])
                     for c in CONFIG_NAMES]
            ax.bar(x + j * width, vals, width, label=f"{cls}-F1")
        ax.set_xticks(x + width / 2)
        ax.set_xticklabels(CONFIG_NAMES, rotation=15)
        ax.set_title(f"Per-class F1 @ {fs}B")
        ax.legend()
    fig.tight_layout()
    fig.savefig(os.path.join(out_dir, "chart_per_class_f1.png"), dpi=150)
    plt.close(fig)

    # 3. confusion matrices (seed 0)
    fig, axes = plt.subplots(len(sizes), len(CONFIG_NAMES), figsize=(4 * len(CONFIG_NAMES), 4 * len(sizes)))
    axes = np.atleast_2d(axes)
    for i, fs in enumerate(sizes):
        for j, config in enumerate(CONFIG_NAMES):
            cm = np.array(results[fs]["ablation"]["configs"][config]["runs"][0]["confusion_matrix"])
            ax = axes[i][j]
            ax.imshow(cm, cmap="Blues")
            ax.set_xticks(range(len(LABELS))); ax.set_xticklabels(LABELS)
            ax.set_yticks(range(len(LABELS))); ax.set_yticklabels(LABELS)
            ax.set_title(f"{config} @ {fs}B")
            for r in range(len(LABELS)):
                for c in range(len(LABELS)):
                    ax.text(c, r, cm[r, c], ha="center", va="center")
    fig.tight_layout()
    fig.savefig(os.path.join(out_dir, "chart_confusion_matrices.png"), dpi=150)
    plt.close(fig)

    # 4. reliability diagrams (before/after, per config, 512B)
    if 512 in results:
        fig, axes = plt.subplots(1, len(CONFIG_NAMES), figsize=(5 * len(CONFIG_NAMES), 4))
        axes = np.atleast_1d(axes)
        for ax, config in zip(axes, CONFIG_NAMES):
            c2 = results[512]["c2"][config]
            ax.plot([0, 1], [0, 1], "k--", label="perfect")
            for key, marker in (("reliability_before", "o"), ("reliability_after", "s")):
                pts = c2[key]
                ax.plot([p["conf"] for p in pts], [p["acc"] for p in pts], marker=marker,
                         label=key.replace("reliability_", ""))
            ax.set_title(f"{config} @ 512B (ECE {c2['ece_before']:.3f}→{c2['ece_after']:.3f})")
            ax.set_xlabel("confidence"); ax.set_ylabel("accuracy")
            ax.legend()
        fig.tight_layout()
        fig.savefig(os.path.join(out_dir, "chart_reliability.png"), dpi=150)
        plt.close(fig)

    # 5. deletion curve
    dc = results.get(DELETION_FRAGMENT_SIZE, {}).get("deletion_curve")
    if dc:
        fig, ax = plt.subplots(figsize=(6, 4))
        ax.plot(dc["fractions"], dc["ig"], marker="o", label="IG-ranked deletion")
        ax.plot(dc["fractions"], dc["random"], marker="s", label="random deletion")
        ax.set_xlabel("fraction of bytes deleted"); ax.set_ylabel("accuracy")
        ax.set_title(f"Deletion curve, +C1 @ {DELETION_FRAGMENT_SIZE}B (n={dc['n_fragments']})")
        ax.legend()
        fig.tight_layout()
        fig.savefig(os.path.join(out_dir, "chart_deletion_curve.png"), dpi=150)
        plt.close(fig)

    # 6. abstention operating curves + 7. ID/OOD confidence histograms
    for fs in sizes:
        fig, ax = plt.subplots(figsize=(6, 4))
        for config in CONFIG_NAMES:
            c4 = results[fs]["c4"][config]
            xs = [r["id_kept_fraction"] for r in c4["rejection_curve"]]
            ys = [r["accuracy_on_kept"] for r in c4["rejection_curve"]]
            ax.plot(xs, ys, marker="o", label=f"{config} (AUROC={c4['auroc']:.3f})")
        ax.set_xlabel("ID fragments kept"); ax.set_ylabel("accuracy on kept")
        ax.set_title(f"Abstention operating curve @ {fs}B")
        ax.legend()
        fig.tight_layout()
        fig.savefig(os.path.join(out_dir, f"chart_rejection_{fs}B.png"), dpi=150)
        plt.close(fig)

        fig, axes = plt.subplots(1, len(CONFIG_NAMES), figsize=(5 * len(CONFIG_NAMES), 3.5))
        axes = np.atleast_1d(axes)
        for ax, config in zip(axes, CONFIG_NAMES):
            c4 = results[fs]["c4"][config]
            ax.hist(c4["id_conf"], bins=20, alpha=0.6, label="ID (doc/docx)", density=True)
            ax.hist(c4["ood_conf"], bins=20, alpha=0.6, label="OOD (pdf/jpeg/random)", density=True)
            ax.set_title(f"{config} @ {fs}B")
            ax.set_xlabel("MSP confidence")
            ax.legend()
        fig.tight_layout()
        fig.savefig(os.path.join(out_dir, f"chart_conf_hist_{fs}B.png"), dpi=150)
        plt.close(fig)


# ----------------------------------------------------------------- excel ---

def write_excel(results: dict, out_dir: str):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title, header, rows):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(row)
        return ws

    # Main results / ablation levels
    rows = []
    for fs, r in results.items():
        for config in CONFIG_NAMES:
            cfg = r["ablation"]["configs"][config]
            doc_f1 = np.mean([x["per_class"]["doc"]["f1"] for x in cfg["runs"]])
            acc = np.mean([x["micro_accuracy"] for x in cfg["runs"]])
            rows.append([f"{fs}B", config, round(cfg["macro_f1_mean"], 4), round(cfg["macro_f1_std"], 4),
                          round(doc_f1, 4), round(cfg["docx_f1_mean"], 4), round(cfg["docx_f1_std"], 4),
                          round(float(acc), 4), cfg["n_params"]])
    sheet("Main Results", ["Fragment size", "Config", "macro-F1 mean", "macro-F1 std",
                             "doc-F1 mean", "docx-F1 mean", "docx-F1 std", "accuracy", "params"], rows)

    # Significance
    rows = []
    for fs, r in results.items():
        sig = r["ablation"]["significance_test"]
        rows.append([f"{fs}B", sig.get("test", "n/a"), sig.get("p_value"), sig.get("effect_size"),
                      sig.get("significant"), sig.get("error", "")])
    sheet("Significance", ["Fragment size", "Test", "p-value", "Effect size", "Significant", "Note"], rows)

    # Per-class (seed 0)
    rows = []
    for fs, r in results.items():
        for config in CONFIG_NAMES:
            run0 = r["ablation"]["configs"][config]["runs"][0]
            for cls in LABELS:
                m = run0["per_class"][cls]
                rows.append([f"{fs}B", config, cls, round(m["precision"], 4), round(m["recall"], 4),
                              round(m["f1"], 4), m["support"]])
    sheet("Per-Class", ["Fragment size", "Config", "Class", "Precision", "Recall", "F1", "Support"], rows)

    # C2 calibration
    rows = []
    for fs, r in results.items():
        for config, c2 in r["c2"].items():
            rows.append([f"{fs}B", config, round(c2["ece_before"], 4), round(c2["ece_after"], 4),
                          round(c2["temperature"], 4)])
    sheet("C2 Calibration", ["Fragment size", "Config", "ECE before", "ECE after", "Temperature"], rows)

    # C2 deletion curve
    dc = results.get(DELETION_FRAGMENT_SIZE, {}).get("deletion_curve")
    if dc:
        rows = [[f, round(ig, 4), round(rnd, 4)] for f, ig, rnd in
                 zip(dc["fractions"], dc["ig"], dc["random"])]
        sheet("C2 Deletion Curve", ["Fraction deleted", "Accuracy (IG-ranked)", "Accuracy (random)"], rows)

    # C4 open-set
    rows = []
    for fs, r in results.items():
        for config, c4 in r["c4"].items():
            rows.append([f"{fs}B", config, round(c4["auroc"], 4), round(c4["id_conf_mean"], 4),
                          round(c4["ood_conf_mean"], 4), c4["n_id"], c4["n_ood"]])
    sheet("C4 Open-Set", ["Fragment size", "Config", "AUROC", "ID conf mean", "OOD conf mean",
                            "n ID", "n OOD"], rows)

    rows = []
    for fs, r in results.items():
        for config, c4 in r["c4"].items():
            for row in c4["rejection_curve"]:
                rows.append([f"{fs}B", config, row["threshold"], round(row["id_kept_fraction"], 4),
                              round(row["accuracy_on_kept"], 4) if row["accuracy_on_kept"] is not None else None,
                              round(row["ood_rejected_fraction"], 4)])
    sheet("C4 Rejection Sweep", ["Fragment size", "Config", "Threshold", "ID kept frac",
                                   "Acc on kept", "OOD rejected frac"], rows)

    # Cost
    rows = []
    for fs, r in results.items():
        for config, cost in r["cost"].items():
            rows.append([f"{fs}B", config, cost["n_params"], cost["flops_per_fragment"],
                          round(cost["throughput_frags_per_sec"], 1)])
    sheet("Cost", ["Fragment size", "Config", "Params", "Est. FLOPs/fragment", "Throughput (frags/s)"], rows)

    # Dataset stats + tuning
    rows = []
    for fs, r in results.items():
        for row in r["corpus_stats"]:
            rows.append([f"{fs}B", row["split"], row["doc"], row["docx"]])
    sheet("Dataset", ["Fragment size", "Split", "doc fragments", "docx fragments"], rows)

    tuning = next(iter(results.values()))["tuning"]
    sheet("Hyperparameters", ["lr", "balance strategy", "epochs", "seeds", "note"],
           [[tuning["lr"], tuning["balance"], tuning["epochs"], str(SEEDS),
             "selected by max min(doc-F1, docx-F1) on val — see hypertune.json"]])

    # Charts sheet with embedded images
    ws = wb.create_sheet("Charts")
    row_anchor = 1
    for png in sorted(p for p in os.listdir(out_dir) if p.startswith("chart_") and p.endswith(".png")):
        try:
            img = XLImage(os.path.join(out_dir, png))
            img.anchor = f"A{row_anchor}"
            ws.add_image(img)
            row_anchor += 24
        except Exception as e:
            ws.cell(row=row_anchor, column=1, value=f"could not embed {png}: {e}")
            row_anchor += 2

    del wb["Sheet"]
    wb.save(os.path.join(out_dir, "FULL_RESULTS.xlsx"))


# ---------------------------------------------------------------- report ---

def write_report(results: dict):
    L = ["# FULL PIPELINE REPORT — C1 + C2 + C3 + C4 (all ablation levels)\n"]
    tuning = next(iter(results.values()))["tuning"]
    L.append(f"Hyperparameters (tuned for balanced doc/docx, objective max min-class-F1): "
             f"lr={tuning['lr']}, balance={tuning['balance']}, epochs={tuning['epochs']}, "
             f"seeds={SEEDS}. Moderate scale; custom doc/docx corpus (FFT-75 not sourced).\n")

    L.append("## C1 — Ablation ladder\n")
    for fs, r in results.items():
        L.append(f"### {fs}B\n")
        L.append("| Config | macro-F1 | docx-F1 |")
        L.append("|---|---|---|")
        for config in CONFIG_NAMES:
            cfg = r["ablation"]["configs"][config]
            L.append(f"| {config} | {cfg['macro_f1_mean']:.3f}±{cfg['macro_f1_std']:.3f} "
                      f"| {cfg['docx_f1_mean']:.3f}±{cfg['docx_f1_std']:.3f} |")
        sig = r["ablation"]["significance_test"]
        if "p_value" in sig:
            L.append(f"\nPaired t-test (backbone vs +C1, macro-F1): p={sig['p_value']:.4f}, "
                      f"effect={sig['effect_size']:+.4f}, significant={sig['significant']}\n")
        else:
            L.append(f"\nSignificance: not computed ({sig.get('error')})\n")
    L.append("![macro-F1 ablation](chart_macro_f1_ablation.png)")
    L.append("![per-class F1](chart_per_class_f1.png)")
    L.append("![confusion matrices](chart_confusion_matrices.png)\n")

    L.append("## C2 — Calibration + attribution faithfulness\n")
    L.append("| Fragment size | Config | ECE before | ECE after | T |")
    L.append("|---|---|---|---|---|")
    for fs, r in results.items():
        for config, c2 in r["c2"].items():
            L.append(f"| {fs}B | {config} | {c2['ece_before']:.4f} | {c2['ece_after']:.4f} "
                      f"| {c2['temperature']:.3f} |")
    L.append("\n![reliability](chart_reliability.png)")
    L.append("![deletion curve](chart_deletion_curve.png)\n")

    L.append("## C3 — Protocol tables\n")
    L.append("Full per-class/per-seed data in FULL_RESULTS.xlsx (sheets: Main Results, Per-Class, "
             "Significance, Dataset).\n")

    L.append("## C4 — Open-set abstention\n")
    L.append("| Fragment size | Config | AUROC | ID conf | OOD conf |")
    L.append("|---|---|---|---|---|")
    for fs, r in results.items():
        for config, c4 in r["c4"].items():
            L.append(f"| {fs}B | {config} | {c4['auroc']:.3f} | {c4['id_conf_mean']:.3f} "
                      f"| {c4['ood_conf_mean']:.3f} |")
    for fs in results:
        L.append(f"\n![rejection {fs}B](chart_rejection_{fs}B.png)")
        L.append(f"![confidence histograms {fs}B](chart_conf_hist_{fs}B.png)")
    L.append("")

    L.append("## Computational cost\n")
    L.append("| Fragment size | Config | params | FLOPs/frag (est.) | frags/s |")
    L.append("|---|---|---|---|---|")
    for fs, r in results.items():
        for config, cost in r["cost"].items():
            L.append(f"| {fs}B | {config} | {cost['n_params']:,} | {cost['flops_per_fragment']:,} "
                      f"| {cost['throughput_frags_per_sec']:.1f} |")
    L.append("")

    with open(os.path.join(OUT_DIR, "FULL_REPORT.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(L))


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    device = "cuda" if torch.cuda.is_available() else "cpu"
    tuning = load_tuning()

    results = {}
    for fragment_size in (512, 4096):
        results[fragment_size] = run_size(fragment_size, device, tuning)
        with open(os.path.join(OUT_DIR, f"full_pipeline_{fragment_size}B.json"), "w") as f:
            json.dump(results[fragment_size], f, indent=2)

    make_charts(results, OUT_DIR)
    write_excel(results, OUT_DIR)
    write_report(results)
    print(f"\nDone: {OUT_DIR}/FULL_REPORT.md, FULL_RESULTS.xlsx, chart_*.png", flush=True)


if __name__ == "__main__":
    main()
